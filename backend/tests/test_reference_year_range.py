"""Batas tahun untuk pengenalan sitasi, dengan dan tanpa koma."""

import pytest
from docx import Document

from app.services.docx_parser import DocxParser
from app.services.reference_validator import ReferenceValidator
from app.services.schema_rules import get_pkm_kc_proposal_rules


@pytest.fixture
def make_validator(tmp_path):
    def make(text, current_year=2026):
        doc = Document()
        doc.add_heading("BAB 1. PENDAHULUAN", level=1)
        doc.add_paragraph(text)
        doc.add_heading("DAFTAR PUSTAKA", level=1)
        doc.add_paragraph(
            "Sari, A. 2020. Judul referensi contoh untuk pemeriksaan sitasi."
        )
        path = tmp_path / "citation_year_range.docx"
        doc.save(path)
        return ReferenceValidator(
            DocxParser(path),
            get_pkm_kc_proposal_rules(),
            current_year=current_year,
            minimum_recommended_recent=0,
        )

    return make


@pytest.mark.parametrize("separator", [", ", " "])
@pytest.mark.parametrize("year", ["2000", "2020", "2026", "2026a"])
def test_citations_in_year_range_are_detected(make_validator, separator, year):
    raw = f"(Sari{separator}{year})"
    citations = make_validator(raw)._extract_in_text_citations()

    assert len(citations) == 1
    assert citations[0].author == "Sari"
    assert citations[0].year == int(year[:4])
    assert citations[0].raw_text == raw
    assert citations[0].has_format_error is (separator == " ")


@pytest.mark.parametrize("separator", [", ", " "])
@pytest.mark.parametrize("year", ["1978", "1999", "2027", "2027a", "6875"])
def test_out_of_range_years_create_no_citation_findings(
    make_validator, separator, year
):
    # Tahun acuan yang lebih baru tidak memperlebar batas tetap 2000–2026.
    result = make_validator(f"(Sari{separator}{year})", current_year=2030).check()

    assert result.in_text_citations == []
    assert result.format_issues == []
    assert result.balance_findings == []
    assert result.status == "pass"


def test_rf_outside_year_range_is_ignored_without_losing_real_citation(make_validator):
    result = make_validator(
        "Hasil (Rf 0,6875) dan (Rf 0,7125) dijelaskan oleh (Sari, 2020)."
    ).check()

    assert [(c.author, c.year) for c in result.in_text_citations] == [("Sari", 2020)]
    assert result.format_issues == []
    assert result.balance_findings == []
    assert result.status == "pass"
