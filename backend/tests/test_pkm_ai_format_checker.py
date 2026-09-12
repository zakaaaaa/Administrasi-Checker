"""
Test suite untuk PkmAiFormatChecker.

Cara jalankan:
    python3 -m pytest tests/test_pkm_ai_format_checker.py -v
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt

from app.services.docx_parser import DocxParser
from app.services.pkm_ai_format_checker import PkmAiFormatChecker
from app.services.schema_rules import get_pkm_ai_proposal_rules

# ============================================================================
# Test: judul yang dibungkus >1 paragraf tidak dikira baris penulis
# ============================================================================
#
# REGRESI lapangan (Leily, PKM-RSH artikel ilmiah): judul artikel panjang
# ditulis 2 paragraf Word — baris pertama style Heading, baris lanjutan
# diketik manual di paragraf baru (tetap ALL CAPS + bold + ~12pt, beda dari
# baris penulis 10pt). Checker lama cuma mengambil paragraf non-kosong
# PERTAMA sebagai judul dan menganggap SEMUA sisanya (sampai heading
# Abstrak) sebagai baris penulis/institusi — baris lanjutan judul lolos
# sebagai "baris penulis" dan divonis salah ukuran font (bukan 10pt) +
# harus normal (bukan bold), padahal itu masih bagian dari judul.


def _build_doc_with_two_paragraph_title(out_path: Path) -> None:
    doc = Document()

    p0 = doc.add_paragraph()
    r0 = p0.add_run("JUDUL ARTIKEL BARIS PERTAMA")
    r0.bold = True
    r0.font.size = Pt(12)

    p1 = doc.add_paragraph()
    r1 = p1.add_run("LANJUTAN JUDUL BARIS KEDUA")
    r1.bold = True
    r1.font.size = Pt(12)

    p2 = doc.add_paragraph()
    r2 = p2.add_run("Budi Santoso, Ani Wijaya")
    r2.font.size = Pt(10)

    doc.add_paragraph("Abstrak")
    doc.add_paragraph(
        "Isi abstrak yang cukup panjang untuk pengujian dan mencapai batas "
        "kata minimum yang wajar untuk sebuah abstrak artikel ilmiah PKM "
        "supaya validasi jumlah kata tidak mengganggu hasil pengujian lain."
    )
    doc.add_paragraph("Pendahuluan")
    doc.add_paragraph("Isi bab pendahuluan.")
    doc.save(str(out_path))


class TestTitleSpanningMultipleParagraphs(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        path = Path(self.tmp.name) / "two_paragraph_title.docx"
        _build_doc_with_two_paragraph_title(path)
        self.parser = DocxParser(path)
        self.rules = get_pkm_ai_proposal_rules()
        self.checker = PkmAiFormatChecker(self.parser, self.rules)

    def tearDown(self):
        self.tmp.cleanup()

    def test_both_title_paragraphs_collected(self):
        bab1_idx = self.checker._find_core_start_para_index()
        title_paras = self.checker._collect_title_paragraphs(bab1_idx)
        texts = [p.text for p in title_paras]
        self.assertEqual(
            texts,
            ["JUDUL ARTIKEL BARIS PERTAMA", "LANJUTAN JUDUL BARIS KEDUA"],
        )

    def test_title_word_count_includes_both_lines(self):
        result = self.checker.check()
        self.assertEqual(result.checks["title_format"].detail["word_count"], 8)

    def test_second_title_line_not_flagged_as_author(self):
        result = self.checker.check()
        joined = " ".join(i.issue for i in result.checks["author_font"].issues)
        self.assertNotIn("LANJUTAN JUDUL", joined)

    def test_real_author_line_still_checked(self):
        """Baris penulis sungguhan (bukan judul) tetap diperiksa — di sini
        fontnya bukan Times New Roman jadi tetap harus ke-flag."""
        result = self.checker.check()
        joined = " ".join(i.issue for i in result.checks["author_font"].issues)
        self.assertIn("Budi Santoso", joined)


if __name__ == "__main__":
    unittest.main(verbosity=2)
