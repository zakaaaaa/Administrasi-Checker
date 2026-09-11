"""
Label penomoran otomatis Word pada judul bab.

Kasus nyata (Laporan Kemajuan PKM-RE): judul bab diketik "PENDAHULUAN" saja,
"BAB 1." di depannya dicetak Word dari numbering.xml yang diikat ke style
Heading 1 (<w:lvlText w:val="BAB %1."/>). Checker membaca teks ketikan saja,
jadi keenam BAB divonis "Tidak Ditemukan" — lalu lembar fisik ("BAB 1 tidak
ditemukan") dan penomoran halaman (halaman inti dianggap zona romawi) ikut
salah vonis.

Cara jalankan:
    python3 -m unittest tests.test_list_numbering_labels -v
"""

import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from app.services.docx_parser import DocxParser
from app.services.physical_sheet_counter import PhysicalSheetCounter
from app.services.schema_rules import get_pkm_laporan_kemajuan_rules
from app.services.structure_checker import StructureChecker

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _numbering_xml(lvl0_text: str = "BAB %1.", lvl0_fmt: str = "decimal") -> str:
    """abstractNum multilevel seperti template kampus: level 0 terikat
    Heading1 ("BAB %1."), level 1 terikat Heading2 ("%1.%2")."""
    return (
        f'<w:numbering xmlns:w="{W_NS}">'
        '<w:abstractNum w:abstractNumId="0">'
        '<w:multiLevelType w:val="multilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/>'
        f'<w:numFmt w:val="{lvl0_fmt}"/><w:pStyle w:val="Heading1"/>'
        f'<w:lvlText w:val="{lvl0_text}"/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/>'
        '<w:numFmt w:val="decimal"/><w:pStyle w:val="Heading2"/>'
        '<w:lvlText w:val="%1.%2"/></w:lvl>'
        "</w:abstractNum>"
        '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        "</w:numbering>"
    )


def _p(text: str, style: str = "", numpr: str = "", rpr: str = "") -> str:
    ppr = ""
    if style or numpr or rpr:
        ppr = (
            "<w:pPr>"
            + (f'<w:pStyle w:val="{style}"/>' if style else "")
            + numpr
            + (f"<w:rPr>{rpr}</w:rPr>" if rpr else "")
            + "</w:pPr>"
        )
    return f"<w:p>{ppr}<w:r><w:t>{text}</w:t></w:r></w:p>"


# numId 0 = penomoran dimatikan untuk paragraf ini (DAFTAR PUSTAKA, LAMPIRAN).
_NO_NUM = '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>'
_SECT = '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>'


def _docx(body_xml: str, numbering_xml: str) -> Path:
    """Dokumen python-docx dengan body, numbering.xml, dan style Heading1/2
    yang bernomor — Heading2 mewarisi numId dari Heading1 lewat basedOn,
    persis dokumen yang dilaporkan."""
    base = Path(tempfile.mkdtemp()) / "base.docx"
    out = Path(tempfile.mkdtemp()) / "numbered.docx"
    Document().save(str(base))
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r"<w:body>.*</w:body>", f"<w:body>{body_xml}</w:body>", xml, flags=re.S)
                data = xml.encode("utf-8")
            elif item.filename == "word/numbering.xml":
                data = numbering_xml.encode("utf-8")
            elif item.filename == "word/styles.xml":
                xml = data.decode("utf-8")
                xml = re.sub(
                    r'(w:styleId="Heading1">.*?<w:pPr>)',
                    r'\1<w:numPr><w:numId w:val="1"/></w:numPr>',
                    xml, count=1, flags=re.S,
                )
                xml = re.sub(
                    r'(w:styleId="Heading2">.*?)<w:basedOn w:val="Normal"/>(.*?<w:pPr>)',
                    r'\1<w:basedOn w:val="Heading1"/>\2<w:numPr><w:ilvl w:val="1"/></w:numPr>',
                    xml, count=1, flags=re.S,
                )
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    return out


def _laporan_body(extra_before_dp: str = "", dp_numpr: str = _NO_NUM) -> str:
    return (
        _p("DAFTAR ISI", "Heading1", _NO_NUM)
        + _p("DAFTAR LAMPIRAN", "Heading1", _NO_NUM)
        + _p("PENDAHULUAN", "Heading1")
        + _p("Latar Belakang", "Heading2")
        + _p("Isi latar belakang.")
        + _p("TARGET LUARAN", "Heading1")
        + _p("Luaran Wajib", "Heading2")
        + _p("Luaran Tambahan", "Heading2")
        + _p("METODE PENELITIAN", "Heading1")
        + _p("HASIL YANG DICAPAI", "Heading1")
        + _p("POTENSI HASIL", "Heading1")
        + _p("RENCANA TAHAPAN BERIKUTNYA", "Heading1")
        + extra_before_dp
        + _p("DAFTAR PUSTAKA", "Heading1", dp_numpr)
        + _p("Penulis, A. (2024). Judul.")
        + _p("LAMPIRAN", "Heading1", _NO_NUM)
        + _SECT
    )


def _by_text(parser: DocxParser) -> dict:
    return {p.text: p for p in parser.paragraphs if p.text}


class TestParserListLabel(unittest.TestCase):
    def setUp(self):
        self.parser = DocxParser(_docx(_laporan_body(), _numbering_xml()))
        self.paras = _by_text(self.parser)

    def test_bab_label_from_heading_style_numbering(self):
        self.assertEqual(self.paras["PENDAHULUAN"].list_label, "BAB 1.")
        self.assertEqual(self.paras["TARGET LUARAN"].list_label, "BAB 2.")
        self.assertEqual(self.paras["RENCANA TAHAPAN BERIKUTNYA"].list_label, "BAB 6.")
        self.assertEqual(self.paras["PENDAHULUAN"].heading_text, "BAB 1. PENDAHULUAN")

    def test_sub_level_inherits_num_through_based_on_and_restarts(self):
        self.assertEqual(self.paras["Latar Belakang"].list_label, "1.1")
        self.assertEqual(self.paras["Luaran Wajib"].list_label, "2.1")
        self.assertEqual(self.paras["Luaran Tambahan"].list_label, "2.2")

    def test_num_id_zero_disables_label(self):
        for text in ("DAFTAR ISI", "DAFTAR LAMPIRAN", "DAFTAR PUSTAKA", "LAMPIRAN"):
            self.assertIsNone(self.paras[text].list_label, text)
        self.assertIsNone(self.paras["Isi latar belakang."].list_label)

    def test_typed_text_is_untouched(self):
        # Modul format/kemiripan membaca `text`; label tidak boleh bocor ke sana.
        self.assertEqual(self.paras["PENDAHULUAN"].text, "PENDAHULUAN")

    def test_hidden_paragraph_mark_has_no_label(self):
        body = _p("PENDAHULUAN", "Heading1", rpr="<w:vanish/>") + _SECT
        parser = DocxParser(_docx(body, _numbering_xml()))
        self.assertIsNone(parser.paragraphs[0].list_label)

    def test_bare_section_break_paragraph_is_not_counted(self):
        # Kasus nyata (LapKem Julian): Heading 1 kosong pembawa section break
        # antara halaman romawi dan halaman inti. Daftar Isi Word dan render
        # LibreOffice sama-sama mencetak "BAB 1. PENDAHULUAN" — paragraf itu
        # tidak bernomor dan tidak dicacah.
        section_break = (
            '<w:p><w:pPr><w:pStyle w:val="Heading1"/>'
            '<w:sectPr><w:pgNumType w:fmt="lowerRoman"/></w:sectPr>'
            "</w:pPr></w:p>"
        )
        body = (
            _p("DAFTAR ISI", "Heading1", _NO_NUM)
            + section_break
            + _p("PENDAHULUAN", "Heading1")
            + _p("TARGET LUARAN", "Heading1")
            + _SECT
        )
        parser = DocxParser(_docx(body, _numbering_xml()))
        self.assertIsNone(parser.paragraphs[1].list_label)
        paras = _by_text(parser)
        self.assertEqual(paras["PENDAHULUAN"].list_label, "BAB 1.")
        self.assertEqual(paras["TARGET LUARAN"].list_label, "BAB 2.")

    def test_empty_numbered_paragraph_without_section_break_still_counts(self):
        # Pengecualian di atas sempit: paragraf bernomor kosong biasa tetap
        # dicetak nomornya oleh Word ("BAB 1." tanpa judul) dan ikut dicacah.
        body = _p("", "Heading1") + _p("PENDAHULUAN", "Heading1") + _SECT
        parser = DocxParser(_docx(body, _numbering_xml()))
        self.assertEqual(parser.paragraphs[0].list_label, "BAB 1.")
        self.assertEqual(_by_text(parser)["PENDAHULUAN"].list_label, "BAB 2.")

    def test_roman_format_is_rendered_as_declared(self):
        parser = DocxParser(_docx(_laporan_body(), _numbering_xml(lvl0_fmt="upperRoman")))
        paras = _by_text(parser)
        self.assertEqual(paras["METODE PENELITIAN"].list_label, "BAB III.")

    def test_unknown_format_is_not_guessed(self):
        parser = DocxParser(_docx(_laporan_body(), _numbering_xml(lvl0_fmt="chineseCounting")))
        self.assertIsNone(_by_text(parser)["PENDAHULUAN"].list_label)

    def test_section_boundaries_see_label(self):
        b = self.parser.find_section_boundaries(["BAB 1", "DAFTAR PUSTAKA"], headings_only=True)
        self.assertEqual(b["BAB 1"], self.paras["PENDAHULUAN"].index)
        self.assertEqual(b["DAFTAR PUSTAKA"], self.paras["DAFTAR PUSTAKA"].index)


class TestCheckersSeeNumberedBab(unittest.TestCase):
    def test_structure_finds_every_bab(self):
        parser = DocxParser(_docx(_laporan_body(), _numbering_xml()))
        result = StructureChecker(parser, get_pkm_laporan_kemajuan_rules("PKM-RE")).check()
        missing = [m.rule_name for m in result.missing_required]
        self.assertFalse([m for m in missing if m.startswith("BAB")], missing)
        self.assertEqual(result.format_violations, [])
        bab1 = next(s for s in result.found_sections if s.rule_name.startswith("BAB 1"))
        self.assertEqual(bab1.matched_text, "BAB 1. PENDAHULUAN")
        self.assertEqual(bab1.match_quality, "exact")

    def test_roman_label_still_flagged_as_format_violation(self):
        # Label dibaca dari deklarasi, bukan dianggap benar: "BAB I." tetap salah.
        parser = DocxParser(_docx(_laporan_body(), _numbering_xml(lvl0_fmt="upperRoman")))
        result = StructureChecker(parser, get_pkm_laporan_kemajuan_rules("PKM-RE")).check()
        self.assertTrue(result.format_violations)

    def test_label_only_adds_matches(self):
        # DAFTAR PUSTAKA yang lupa dimatikan penomorannya tercetak
        # "BAB 7. DAFTAR PUSTAKA" — tetap harus dikenali sebagai DAFTAR PUSTAKA.
        parser = DocxParser(_docx(_laporan_body(dp_numpr=""), _numbering_xml()))
        self.assertEqual(_by_text(parser)["DAFTAR PUSTAKA"].list_label, "BAB 7.")
        result = StructureChecker(parser, get_pkm_laporan_kemajuan_rules("PKM-RE")).check()
        self.assertNotIn("DAFTAR PUSTAKA", [m.rule_name for m in result.missing_required])

    def test_physical_sheet_locates_bab1(self):
        parser = DocxParser(_docx(_laporan_body(), _numbering_xml()))
        counter = PhysicalSheetCounter(parser, get_pkm_laporan_kemajuan_rules("PKM-RE"))
        bab1, lampiran = counter._locate_core_paragraphs()
        paras = _by_text(parser)
        self.assertEqual(bab1, paras["PENDAHULUAN"].index)
        self.assertEqual(lampiran, paras["LAMPIRAN"].index)


if __name__ == "__main__":
    unittest.main()
