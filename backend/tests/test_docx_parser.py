"""
Test suite untuk DocxParser — versi unittest (stdlib).

Bisa dijalankan dengan:
    python3 -m unittest tests.test_docx_parser -v
atau (jika pytest tersedia):
    python3 -m pytest tests/test_docx_parser.py -v
"""

import unittest
import tempfile
from pathlib import Path

from app.services.docx_parser import (
    DocxParser,
    ParagraphInfo,
    TableInfo,
    SectionInfo,
    dxa_to_cm,
    half_points_to_pt,
)

SAMPLE_DIR = Path(__file__).parent / "sample_docs"
DUMMY_FILE = SAMPLE_DIR / "dummy_pkm_kc.docx"


class DocxParserTestBase(unittest.TestCase):
    """Base class — load parser sekali untuk efisiensi."""

    @classmethod
    def setUpClass(cls):
        if not DUMMY_FILE.exists():
            raise unittest.SkipTest(
                "Dummy belum di-generate. "
                "Jalankan: python3 tests/build_dummy_docx.py"
            )
        cls.parser = DocxParser(DUMMY_FILE)


# ============================================================================
# Test: konstruktor & validasi input
# ============================================================================


class TestConstructor(unittest.TestCase):
    def test_file_not_found_raises(self):
        with self.assertRaises(FileNotFoundError):
            DocxParser("/path/yang/tidak/ada.docx")

    def test_non_docx_extension_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            fake = Path(tmp) / "bukan.pdf"
            fake.write_bytes(b"hello")
            with self.assertRaises(ValueError):
                DocxParser(fake)


# ============================================================================
# Test: property dasar
# ============================================================================


class TestBasicProperties(DocxParserTestBase):
    def test_paragraphs_extracted(self):
        paras = self.parser.paragraphs
        self.assertIsInstance(paras, list)
        self.assertGreater(len(paras), 0)
        for p in paras:
            self.assertIsInstance(p, ParagraphInfo)
        self.assertTrue(any(p.text.strip() for p in paras))

    def test_paragraph_has_index_text_runs(self):
        for p in self.parser.paragraphs:
            self.assertIsInstance(p.index, int)
            self.assertIsInstance(p.text, str)
            self.assertIsInstance(p.runs, list)

    def test_heading_detection(self):
        bab1 = [p for p in self.parser.paragraphs if "BAB 1" in p.text.upper()]
        self.assertGreaterEqual(len(bab1), 1, "BAB 1 tidak ditemukan")
        heading_paras = [p for p in bab1 if p.is_heading]
        self.assertGreaterEqual(
            len(heading_paras), 1,
            "BAB 1 ada tapi tidak terdeteksi sebagai heading"
        )

    def test_tables_extracted(self):
        tables = self.parser.tables
        self.assertIsInstance(tables, list)
        self.assertGreaterEqual(len(tables), 2, "Dummy punya 2 tabel")
        for t in tables:
            self.assertIsInstance(t, TableInfo)
            self.assertGreater(t.rows, 0)
            self.assertGreater(t.cols, 0)
            self.assertEqual(len(t.cells), t.rows * t.cols)

    def test_table_header_texts(self):
        rab_tables = [
            t for t in self.parser.tables
            if any("Jenis Pengeluaran" in h for h in t.header_texts)
        ]
        self.assertGreaterEqual(len(rab_tables), 1, "Tabel RAB tidak terdeteksi")

    def test_sections_extracted(self):
        sections = self.parser.sections
        self.assertIsInstance(sections, list)
        self.assertGreaterEqual(len(sections), 1)
        for s in sections:
            self.assertIsInstance(s, SectionInfo)

    def test_section_margins_match_pkm_rules(self):
        """
        Margin di-baca langsung dari XML <w:pgMar>, jadi nilai exact.
        Dummy generate dengan Cm(4) untuk kiri dan Cm(3) untuk lainnya.
        Cm(4) → 4 × 567 ≈ 2268 DXA → /1440 × 2.54 = 4.0 cm ✓
        """
        s0 = self.parser.sections[0]
        self.assertEqual(dxa_to_cm(s0.margin_left_dxa), 4.0)
        self.assertEqual(dxa_to_cm(s0.margin_right_dxa), 3.0)
        self.assertEqual(dxa_to_cm(s0.margin_top_dxa), 3.0)
        self.assertEqual(dxa_to_cm(s0.margin_bottom_dxa), 3.0)

    def test_section_margins_not_zero(self):
        """
        Regression test untuk bug v1 di dokumen real: margin terbaca 0
        karena konversi Twips→EMU→DXA via wrapper python-docx tidak konsisten.
        Setelah fix (baca langsung dari XML), tidak ada margin yang 0
        kecuali memang di-set 0 di dokumen.
        """
        for sec in self.parser.sections:
            # Dummy kita set margin > 0 untuk semua side
            self.assertGreater(sec.margin_left_dxa, 0,
                f"Section #{sec.index} margin_left = 0, kemungkinan bug parsing")
            self.assertGreater(sec.margin_right_dxa, 0,
                f"Section #{sec.index} margin_right = 0")
            self.assertGreater(sec.margin_top_dxa, 0,
                f"Section #{sec.index} margin_top = 0")
            self.assertGreater(sec.margin_bottom_dxa, 0,
                f"Section #{sec.index} margin_bottom = 0")


# ============================================================================
# Test: header/footer XML
# ============================================================================


class TestRawXMLAccess(DocxParserTestBase):
    def test_header_footer_xmls_accessible(self):
        self.assertIsInstance(self.parser.header_xmls, dict)
        self.assertIsInstance(self.parser.footer_xmls, dict)

    def test_document_xml_loads(self):
        root = self.parser.document_xml
        self.assertIsNotNone(root)
        self.assertTrue(root.tag.endswith("}document"))


# ============================================================================
# Test: find_section_boundaries
# ============================================================================


class TestSectionBoundaries(DocxParserTestBase):
    def test_default_matches_first_occurrence(self):
        """
        Default (headings_only=False): match kemunculan pertama, yang BISA berupa
        baris di Daftar Isi. Ini adalah perilaku lama yang dipertahankan.
        """
        b = self.parser.find_section_boundaries(
            ["DAFTAR ISI", "BAB 1", "DAFTAR PUSTAKA", "LAMPIRAN"]
        )
        self.assertIsNotNone(b["DAFTAR ISI"])
        self.assertIsNotNone(b["BAB 1"])
        self.assertIsNotNone(b["DAFTAR PUSTAKA"])
        self.assertIsNotNone(b["LAMPIRAN"])
        # Note: di dummy, "BAB 1" di-match ke baris ToC (paragraf #1),
        # bukan ke heading BAB 1 asli. Itu memang perilaku default.

    def test_headings_only_skips_toc_entries(self):
        """
        headings_only=True: hanya match paragraf yang is_heading.
        Baris di Daftar Isi (yang bukan heading) di-skip, sehingga
        "BAB 1" akan ketemu di heading aslinya, bukan di entri ToC.
        """
        b_default = self.parser.find_section_boundaries(["BAB 1"])
        b_headings = self.parser.find_section_boundaries(
            ["BAB 1"], headings_only=True
        )
        self.assertIsNotNone(b_default["BAB 1"])
        self.assertIsNotNone(b_headings["BAB 1"])
        # Heading asli pasti BELAKANG (lebih besar indeksnya) daripada entri ToC.
        self.assertGreater(
            b_headings["BAB 1"], b_default["BAB 1"],
            "Heading asli BAB 1 seharusnya muncul setelah baris BAB 1 di ToC"
        )
        # Verifikasi yang ketemu memang heading
        bab1_para = self.parser.paragraphs[b_headings["BAB 1"]]
        self.assertTrue(bab1_para.is_heading)
        self.assertIn("PENDAHULUAN", bab1_para.text.upper())

    def test_headings_only_full_proposal_order(self):
        """
        Dengan headings_only=True, urutan section dalam proposal harus benar:
        DAFTAR ISI < BAB 1 < BAB 2 < BAB 3 < BAB 4 < DAFTAR PUSTAKA < LAMPIRAN.
        """
        b = self.parser.find_section_boundaries(
            ["DAFTAR ISI", "BAB 1", "BAB 2", "BAB 3", "BAB 4",
             "DAFTAR PUSTAKA", "LAMPIRAN"],
            headings_only=True,
        )
        # Semua harus ditemukan
        for name, idx in b.items():
            self.assertIsNotNone(idx, f"Section {name!r} tidak ditemukan")
        # Urutan strictly ascending
        order = ["DAFTAR ISI", "BAB 1", "BAB 2", "BAB 3", "BAB 4",
                 "DAFTAR PUSTAKA", "LAMPIRAN"]
        for prev, nxt in zip(order, order[1:]):
            self.assertLess(
                b[prev], b[nxt],
                f"Urutan salah: {prev}({b[prev]}) seharusnya sebelum {nxt}({b[nxt]})"
            )

    def test_missing_returns_none(self):
        b = self.parser.find_section_boundaries(["BAGIAN_FIKTIF_TIDAK_ADA"])
        self.assertIsNone(b["BAGIAN_FIKTIF_TIDAK_ADA"])

    def test_missing_with_headings_only(self):
        """Section yang tidak ada → None, baik dengan flag maupun tanpa."""
        b = self.parser.find_section_boundaries(
            ["BAGIAN_FIKTIF"], headings_only=True
        )
        self.assertIsNone(b["BAGIAN_FIKTIF"])


# ============================================================================
# Test: iter_runs
# ============================================================================


class TestIterRuns(DocxParserTestBase):
    def test_iter_runs_yields_data(self):
        count = 0
        found_tnr = False
        for _, run in self.parser.iter_runs():
            count += 1
            if run.font_name and "Times New Roman" in run.font_name:
                found_tnr = True
        self.assertGreater(count, 0)
        self.assertTrue(found_tnr, "Run dengan TNR tidak ditemukan")


# ============================================================================
# Test: images
# ============================================================================


class TestImages(DocxParserTestBase):
    def test_images_property(self):
        images = self.parser.images
        self.assertIsInstance(images, list)
        # Dummy tidak ada gambar — kosong wajar


# ============================================================================
# Test: konversi unit
# ============================================================================


class TestUnitConversion(unittest.TestCase):
    def test_dxa_to_cm(self):
        self.assertAlmostEqual(dxa_to_cm(11906), 21.0, delta=0.01)
        self.assertIsNone(dxa_to_cm(None))
        self.assertAlmostEqual(dxa_to_cm(1440), 2.54, delta=0.001)

    def test_half_points_to_pt(self):
        self.assertEqual(half_points_to_pt(24), 12.0)
        self.assertEqual(half_points_to_pt("24"), 12.0)
        self.assertIsNone(half_points_to_pt(None))
        self.assertIsNone(half_points_to_pt("invalid"))


# ============================================================================
# Test: summary
# ============================================================================


class TestSummary(DocxParserTestBase):
    def test_summary_returns_dict(self):
        s = self.parser.summary()
        self.assertIsInstance(s, dict)
        expected_keys = {
            "file", "paragraph_count", "table_count", "section_count",
            "header_part_count", "footer_part_count", "image_count", "warnings",
        }
        self.assertTrue(expected_keys.issubset(set(s.keys())))
        self.assertGreater(s["paragraph_count"], 0)
        self.assertGreaterEqual(s["table_count"], 2)


# ============================================================================
# Test: invalid zip
# ============================================================================


class TestCorruptedFile(unittest.TestCase):
    def test_corrupted_docx_collects_warning(self):
        with tempfile.TemporaryDirectory() as tmp:
            fake = Path(tmp) / "rusak.docx"
            fake.write_bytes(b"ini bukan zip valid")
            parser = DocxParser(fake)
            raw = parser.read_raw_part("word/document.xml")
            self.assertIsNone(raw)
            self.assertGreater(len(parser.warnings), 0)


# ============================================================================
# Test: enumerasi section & warisan header/footer
# ============================================================================
#
# python-docx `doc.sections` hanya mengenumerasi sectPr pada paragraf
# tingkat-body dan sectPr penutup body. Section break yang bersarang di dalam
# content control (<w:sdt>) — lazim pada dokumen ber-Daftar Isi otomatis —
# hilang dari daftar, BESERTA header/footer yang dideklarasikannya.
#
# Kasus lapangan (Laporan Kemajuan PKM-PM): section break yang memulai BAB 1
# bersarang di <w:sdt> dan membawa headerReference berisi nomor halaman arab.
# Karena section itu tak terlihat, checker memvonis "nomor halaman arab tidak
# ditemukan di bagian isi" — padahal di dokumen aslinya jelas tercetak.
#
# Warisan: section tanpa referensi sendiri memakai milik section sebelumnya
# per tipe (ECMA-376 §17.10.1). Banyak dokumen mendeklarasikan nomor halaman
# sekali saja di section pertama.
# ============================================================================

from docx.oxml.ns import qn as _qn  # noqa: E402
from lxml import etree as _etree  # noqa: E402

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _build_body_xml(inner: str) -> str:
    return (
        f'<w:document xmlns:w="{_W}" xmlns:r="{_R}"><w:body>{inner}</w:body></w:document>'
    )


def _sect_pr(header_rid=None, footer_rid=None) -> str:
    refs = ""
    if header_rid:
        refs += f'<w:headerReference w:type="default" r:id="{header_rid}"/>'
    if footer_rid:
        refs += f'<w:footerReference w:type="default" r:id="{footer_rid}"/>'
    return f"<w:sectPr>{refs}<w:pgSz w:w=\"11906\" w:h=\"16838\"/></w:sectPr>"


class _XmlOnlyParser(DocxParser):
    """Parser yang hanya perlu document_xml — melewati pembacaan file."""

    def __init__(self, xml: str):  # noqa: D107  (sengaja tidak panggil super)
        self.warnings = []
        self._document_xml = _etree.fromstring(xml.encode())
        self._sections = None


class TestSectionEnumeration(unittest.TestCase):
    def test_sect_pr_nested_in_sdt_is_counted(self):
        """REGRESI PKM-PM: section break di dalam content control."""
        xml = _build_body_xml(
            f'<w:p><w:pPr>{_sect_pr("rId9")}</w:pPr></w:p>'
            f'<w:sdt><w:sdtContent><w:p><w:pPr>{_sect_pr("rId18")}</w:pPr></w:p>'
            f"</w:sdtContent></w:sdt>"
            f'{_sect_pr()}'
        )
        sections = _XmlOnlyParser(xml).sections
        self.assertEqual(len(sections), 3)
        self.assertEqual(sections[0].header_refs.get("default"), "rId9")
        self.assertEqual(sections[1].header_refs.get("default"), "rId18")

    def test_sect_pr_change_is_ignored(self):
        """<w:sectPrChange> memuat salinan sectPr LAMA — bukan section nyata."""
        xml = _build_body_xml(
            f'<w:p><w:pPr>{_sect_pr("rId9")}'
            f'<w:sectPrChange w:id="1">{_sect_pr("rId99")}</w:sectPrChange>'
            f"</w:pPr></w:p>"
            f'{_sect_pr()}'
        )
        sections = _XmlOnlyParser(xml).sections
        self.assertEqual(len(sections), 2)
        rids = {s.header_refs.get("default") for s in sections}
        self.assertNotIn("rId99", rids)


class TestHeaderFooterInheritance(unittest.TestCase):
    def test_section_without_refs_inherits_previous(self):
        xml = _build_body_xml(
            f'<w:p><w:pPr>{_sect_pr("rId9", "rId11")}</w:pPr></w:p>'
            f'<w:p><w:pPr>{_sect_pr()}</w:pPr></w:p>'
            f'{_sect_pr()}'
        )
        sections = _XmlOnlyParser(xml).sections
        self.assertEqual(len(sections), 3)
        for sec in sections:
            with self.subTest(section=sec.index):
                self.assertEqual(sec.header_refs.get("default"), "rId9")
                self.assertEqual(sec.footer_refs.get("default"), "rId11")

    def test_own_reference_overrides_inherited(self):
        xml = _build_body_xml(
            f'<w:p><w:pPr>{_sect_pr("rId9")}</w:pPr></w:p>'
            f'<w:p><w:pPr>{_sect_pr("rId18")}</w:pPr></w:p>'
            f'<w:p><w:pPr>{_sect_pr()}</w:pPr></w:p>'
            f'{_sect_pr()}'
        )
        sections = _XmlOnlyParser(xml).sections
        self.assertEqual(sections[0].header_refs["default"], "rId9")
        self.assertEqual(sections[1].header_refs["default"], "rId18")
        # Section ke-3 tanpa deklarasi → warisi yang TERBARU, bukan yang pertama.
        self.assertEqual(sections[2].header_refs["default"], "rId18")

    def test_inheritance_does_not_mutate_across_sections(self):
        """Tiap section punya dict sendiri — bukan referensi bersama."""
        xml = _build_body_xml(
            f'<w:p><w:pPr>{_sect_pr("rId9")}</w:pPr></w:p>'
            f'<w:p><w:pPr>{_sect_pr("rId18")}</w:pPr></w:p>'
            f'{_sect_pr()}'
        )
        sections = _XmlOnlyParser(xml).sections
        self.assertIsNot(sections[0].header_refs, sections[1].header_refs)
        self.assertEqual(sections[0].header_refs["default"], "rId9")



# ============================================================================
# Dokumen dengan gambar rusak tetap bisa diperiksa
# ============================================================================
#
# Satu gambar yang CRC-nya rusak membuat python-docx gagal membuka seluruh
# dokumen ("Bad CRC-32 for file 'word/media/image1.png'"). Semua modul crash,
# tak satu pun pengecekan jalan — dan pesan galat itu sempat tampil di UI
# sebagai "Kesalahan Judul Bab (word/media/image1.png)". Padahal teks
# dokumennya utuh dan Word membukanya tanpa masalah.
# ============================================================================

CORRUPT_FILE = SAMPLE_DIR / "corrupt_image.docx"


class TestCorruptMediaTolerance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not CORRUPT_FILE.exists():
            raise unittest.SkipTest("sampel corrupt_image.docx tidak tersedia")

    def test_sample_really_is_corrupt(self):
        """Prasyarat: sampelnya memang rusak, bukan lolos karena kebetulan utuh."""
        import zipfile

        self.assertIsNotNone(zipfile.ZipFile(CORRUPT_FILE).testzip())

    def test_document_opens_despite_corrupt_image(self):
        parser = DocxParser(CORRUPT_FILE)
        self.assertGreater(len(parser.paragraphs), 50)
        self.assertEqual(parser.paragraphs[0].text.strip(), "DAFTAR ISI")

    def test_corrupt_part_is_reported(self):
        parser = DocxParser(CORRUPT_FILE)
        self.assertEqual(len(parser.corrupt_parts), 1)
        self.assertTrue(parser.corrupt_parts[0].startswith("word/media/"))
        self.assertTrue(any("rusak" in w for w in parser.warnings))

    def test_original_path_is_kept(self):
        """Salinan perbaikan dipakai untuk membaca, path asli tetap tercatat."""
        parser = DocxParser(CORRUPT_FILE)
        self.assertEqual(parser.original_file_path, CORRUPT_FILE)
        self.assertNotEqual(parser.file_path, CORRUPT_FILE)

    def test_healthy_document_untouched(self):
        """Dokumen sehat tidak boleh ikut disalin/diubah."""
        parser = DocxParser(DUMMY_FILE)
        self.assertEqual(parser.corrupt_parts, [])
        self.assertEqual(parser.file_path, DUMMY_FILE)



# ============================================================================
# Section yang terperangkap di dalam content control (<w:sdt>)
# ============================================================================
#
# REGRESI lapangan (Lapkem PKM-RSH, Mendeley): judul LAMPIRAN dan seluruh
# lampiran ikut masuk ke content control bibliografi. Paragraf di dalam
# <w:sdt> tidak dienumerasi python-docx → LAMPIRAN divonis tidak ada dan isi
# lampiran terbaca sebagai entri Daftar Pustaka.


def _docx_with_body(body_xml: str) -> Path:
    """Dokumen nyata (python-docx) yang body-nya diganti dengan XML mentah."""
    import re as _re
    import zipfile as _zip
    from docx import Document as _Document

    tmp = Path(tempfile.mkdtemp()) / "sdt.docx"
    base = Path(tempfile.mkdtemp()) / "base.docx"
    _Document().save(str(base))
    with _zip.ZipFile(base) as zin, _zip.ZipFile(tmp, "w", _zip.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                xml = _re.sub(r"<w:body>.*</w:body>", f"<w:body>{body_xml}</w:body>", xml, flags=_re.S)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    return tmp


def _p(text: str, style: str = "") -> str:
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{ppr}<w:r><w:t>{text}</w:t></w:r></w:p>"


def _sdt(inner: str, gallery: str = "") -> str:
    pr = (
        f'<w:sdtPr><w:docPartObj><w:docPartGallery w:val="{gallery}"/></w:docPartObj></w:sdtPr>'
        if gallery
        else '<w:sdtPr><w:tag w:val="MENDELEY_BIBLIOGRAPHY"/></w:sdtPr>'
    )
    return f"<w:sdt>{pr}<w:sdtContent>{inner}</w:sdtContent></w:sdt>"


_SECT = '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>'


class TestManualTableOfContentsBlock(unittest.TestCase):
    """Regresi NeuroRehab: beberapa entri Daftar Isi manual kehilangan leader."""

    _PAGE_BREAK = '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'

    def test_heading_shaped_gaps_between_manual_entries_are_toc(self):
        path = _docx_with_body(
            _p("DAFTAR ISI")
            + _p("BAB 1. PENDAHULUAN ........................ 1")
            + _p("BAB 2. TARGET LUARAN")
            + _p("BAB 3. TAHAP PELAKSANAAN")
            + _p("BAB 4. HASIL YANG DICAPAI ................. 7")
            + self._PAGE_BREAK
            + _p("BAB 1. PENDAHULUAN", "Heading1")
            + _SECT
        )

        paragraphs = DocxParser(path).paragraphs
        by_text = {p.text: p for p in paragraphs if p.text}

        self.assertEqual(by_text["BAB 1. PENDAHULUAN ........................ 1"].toc_evidence, "leader")
        self.assertEqual(by_text["BAB 2. TARGET LUARAN"].toc_evidence, "manual-block")
        self.assertEqual(by_text["BAB 3. TAHAP PELAKSANAAN"].toc_evidence, "manual-block")
        self.assertTrue(by_text["BAB 2. TARGET LUARAN"].is_toc_entry)
        self.assertFalse(by_text["BAB 2. TARGET LUARAN"].is_heading)
        self.assertFalse(by_text["BAB 1. PENDAHULUAN"].is_toc_entry)
        self.assertTrue(by_text["BAB 1. PENDAHULUAN"].is_heading)

    def test_single_leader_is_not_enough_to_infer_manual_block(self):
        path = _docx_with_body(
            _p("DAFTAR ISI")
            + _p("BAB 1. PENDAHULUAN .......... 1")
            + _p("BAB 2. TARGET LUARAN")
            + self._PAGE_BREAK
            + _p("BAB 1. PENDAHULUAN", "Heading1")
            + _SECT
        )

        paragraphs = DocxParser(path).paragraphs
        bab2 = next(p for p in paragraphs if p.text == "BAB 2. TARGET LUARAN")
        self.assertFalse(bab2.is_toc_entry)
        self.assertIsNone(bab2.toc_evidence)


class TestSectionTrappedInSdt(unittest.TestCase):
    def test_lampiran_inside_bibliography_is_hoisted(self):
        path = _docx_with_body(
            _p("DAFTAR PUSTAKA", "Heading1")
            + _sdt(
                _p("Anandi, R.D. (2022) Judul.")
                + _p("Wongso, D. A. (2025) Judul.")
                + _p("LAMPIRAN", "Heading1")
                + _p("Lampiran 1. Penggunaan Dana", "Caption")
            )
            + _p("Lampiran 6.3 Tabel")
            + _SECT
        )
        parser = DocxParser(path)
        texts = [p.text for p in parser.paragraphs]
        self.assertEqual(
            texts,
            ["DAFTAR PUSTAKA", "LAMPIRAN", "Lampiran 1. Penggunaan Dana", "Lampiran 6.3 Tabel"],
        )
        self.assertEqual(parser.hoisted_sdt_elements, 2)
        # Entri pustaka tetap di dalam content control.
        body = parser.document_xml.find("w:body", {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
        sdt_text = "".join(body[1].itertext())
        self.assertIn("Wongso", sdt_text)
        self.assertNotIn("LAMPIRAN", sdt_text)

    def test_document_xml_stays_aligned_with_paragraphs(self):
        path = _docx_with_body(
            _p("DAFTAR PUSTAKA", "Heading1")
            + _sdt(_p("Anandi (2022).") + _p("LAMPIRAN") + _p("isi lampiran"))
            + _SECT
        )
        parser = DocxParser(path)
        body_ps = [c for c in parser.document_xml.find(f"{{{W}}}body") if c.tag == f"{{{W}}}p"]
        self.assertEqual(len(body_ps), len(parser.paragraphs))
        self.assertEqual(["".join(p.itertext()) for p in body_ps], [p.text for p in parser.paragraphs])

    def test_table_of_contents_sdt_untouched(self):
        path = _docx_with_body(
            _sdt(_p("x") + _p("DAFTAR ISI", "Heading1") + _p("BAB 1\t1", "TOC1"), gallery="Table of Contents")
            + _p("BAB 1. PENDAHULUAN", "Heading1")
            + _SECT
        )
        parser = DocxParser(path)
        self.assertEqual([p.text for p in parser.paragraphs], ["BAB 1. PENDAHULUAN"])
        self.assertEqual(parser.hoisted_sdt_elements, 0)

    def test_bibliography_without_heading_untouched(self):
        path = _docx_with_body(
            _p("DAFTAR PUSTAKA", "Heading1")
            + _sdt(_p("Anandi (2022).") + _p("Wongso (2025)."))
            + _p("LAMPIRAN", "Heading1")
            + _SECT
        )
        parser = DocxParser(path)
        self.assertEqual([p.text for p in parser.paragraphs], ["DAFTAR PUSTAKA", "LAMPIRAN"])
        self.assertEqual(parser.hoisted_sdt_elements, 0)

    def test_heading_as_first_child_stays_inside(self):
        """Blok yang memang diawali judulnya sendiri dibiarkan utuh."""
        path = _docx_with_body(
            _sdt(_p("DAFTAR PUSTAKA", "Heading1") + _p("Anandi (2022)."))
            + _p("LAMPIRAN", "Heading1")
            + _SECT
        )
        parser = DocxParser(path)
        self.assertEqual([p.text for p in parser.paragraphs], ["LAMPIRAN"])
        self.assertEqual(parser.hoisted_sdt_elements, 0)


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ============================================================================
# Estimasi halaman: LRPB di baris tabel yang terbelah
# ============================================================================
#
# REGRESI lapangan (Lapkem PKM-PM "Tim Sabar"): baris terakhir Tabel 1
# terbelah ke halaman berikut. Word menulis LRPB tiap kali halaman berpindah
# menurut urutan dokumen — termasuk saat sel kedua KEMBALI ke puncak baris di
# halaman lama, dan di awal paragraf sesudah tabel. Aturan lama menjumlah
# semua LRPB "lead" di tabel sehingga satu belahan terhitung dua halaman;
# Daftar Pustaka (hal. 10) terhitung hal. 11 → "bagian inti melebihi 10".


def _lp(*segs: str, ppr: str = "") -> str:
    """Paragraf dari potongan teks; "|" = lastRenderedPageBreak."""
    runs = "".join(
        "<w:r><w:lastRenderedPageBreak/></w:r>" if s == "|" else f"<w:r><w:t>{s}</w:t></w:r>"
        for s in segs
    )
    return f"<w:p>{ppr}{runs}</w:p>"


def _tbl(*rows: list[str]) -> str:
    trs = "".join(
        "<w:tr>" + "".join(f"<w:tc>{c}</w:tc>" for c in cells) + "</w:tr>" for cells in rows
    )
    return f"<w:tbl>{trs}</w:tbl>"


def _pages(body_xml: str) -> dict[str, int]:
    parser = DocxParser(_docx_with_body(body_xml + _SECT))
    return {p.text: parser.estimate_physical_page(p.index) for p in parser.paragraphs}


class TestSplitTableRowPageEstimate(unittest.TestCase):
    def test_split_last_row_counts_one_page(self):
        """Pola Tim Sabar: sel 0 terbelah, sel 1 kembali ke atas lalu
        terbelah, paragraf sesudah tabel diawali LRPB artefak."""
        pages = _pages(
            _lp("Awal")
            + _tbl(
                [_lp("h1"), _lp("h2")],
                [_lp("kiri atas", "|", "kiri bawah"), _lp("|", "kanan atas", "|", "kanan bawah")],
            )
            + _lp("|", "Sesudah tabel")
            + _lp("|", "Halaman berikut")
        )
        self.assertEqual(pages["Awal"], 1)
        self.assertEqual(pages["Sesudah tabel"], 2)
        self.assertEqual(pages["Halaman berikut"], 3)

    def test_split_row_many_columns_is_one_break(self):
        """Tiap sel yang kembali ke puncak baris dapat LRPB di awal —
        empat kolom tidak boleh jadi empat halaman."""
        cell = _lp("|", "atas", "|", "bawah")
        pages = _pages(
            _lp("Awal")
            + _tbl([_lp("atas", "|", "bawah"), cell, cell, cell], [_lp("|", "baris lanjut"), _lp("x"), _lp("y"), _lp("z")])
            + _lp("Sesudah tabel")
        )
        self.assertEqual(pages["Sesudah tabel"], 2)

    def test_row_starting_new_page_still_counted(self):
        """Pemutus antarbaris biasa (LRPB di awal sel pertama) tetap dihitung."""
        pages = _pages(
            _lp("Awal")
            + _tbl([_lp("r1"), _lp("r1b")], [_lp("|", "r2"), _lp("r2b")], [_lp("|", "r3"), _lp("r3b")])
            + _lp("Sesudah tabel")
        )
        self.assertEqual(pages["Sesudah tabel"], 3)

    def test_forced_break_after_split_row_kept(self):
        """Paragraf sesudah baris terbelah yang MEMANG memaksa halaman baru
        (page break manual / pageBreakBefore) tidak dianggap artefak."""
        split = _tbl([_lp("kiri", "|", "bawah"), _lp("|", "kanan", "|", "bawah")])
        manual = (
            '<w:p><w:r><w:br w:type="page"/></w:r>'
            "<w:r><w:lastRenderedPageBreak/><w:t>BAB 3</w:t></w:r></w:p>"
        )
        pages = _pages(_lp("Awal") + split + manual)
        self.assertEqual(pages["BAB 3"], 3)
        pbb = _lp("|", "BAB 4", ppr="<w:pPr><w:pageBreakBefore/></w:pPr>")
        pages = _pages(_lp("Awal") + split + pbb)
        self.assertEqual(pages["BAB 4"], 3)

    def test_paragraph_lrpb_without_table_unchanged(self):
        pages = _pages(_lp("Satu") + _lp("|", "Dua") + _lp("tiga", "|", "empat") + _lp("Lima"))
        self.assertEqual(pages, {"Satu": 1, "Dua": 2, "tigaempat": 2, "Lima": 3})


# ============================================================================
# Test: style judul custom yang namanya bukan "Heading N"
# ============================================================================
#
# REGRESI lapangan (PKM-RE Artikel Ilmiah, template Word bawaan
# penerbit/kampus): section "Pendahuluan"/"Metode"/"Daftar Pustaka" diberi
# style "Article Heading 1" — bukan style bawaan Word ("Heading 1") dan
# tanpa outlineLvl. Cek lama (name.startswith("heading")) melewatkannya
# karena nama dimulai "Article", bukan "Heading", jadi is_heading tetap
# False dan section itu divonis hilang oleh StructureChecker padahal ada.


class TestCustomHeadingStyleName(unittest.TestCase):
    def test_style_name_containing_heading_word_is_detected(self):
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE

        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            style = doc.styles.add_style("Article Heading 1", WD_STYLE_TYPE.PARAGRAPH)
            style.font.bold = True  # bold di STYLE, bukan di run — tidak lolos heuristik bold-per-run
            doc.add_paragraph("Pendahuluan", style=style)
            doc.add_paragraph("Teks biasa tidak boleh ikut ke-flag sebagai heading.")
            path = Path(tmp) / "custom_heading.docx"
            doc.save(str(path))

            parser = DocxParser(path)
            heading_p = next(p for p in parser.paragraphs if p.text == "Pendahuluan")
            body_p = next(p for p in parser.paragraphs if p.text.startswith("Teks biasa"))
            self.assertTrue(heading_p.is_heading)
            self.assertFalse(body_p.is_heading)


# ============================================================================
# Test: line_spacing mode "Exactly"/"At least" tidak dibalik jadi EMU mentah
# ============================================================================
#
# REGRESI lapangan (PKM-RSH artikel ilmiah, front matter judul/penulis):
# python-docx balikin objek Length (EMU) untuk paragraph_format.line_spacing
# kalau line_spacing_rule-nya EXACTLY/AT_LEAST (bukan MULTIPLE) — float(ls)
# lama-lama jadi angka mentah ratusan ribu (mis. 144780.0 untuk "Exactly
# 11.4pt"), lalu dibandingkan checker seolah multiplier ("harus 1.0") dan
# menghasilkan pesan absurd "spasi baris ditemukan 144780.0".


class TestLineSpacingExactModeNotMisread(unittest.TestCase):
    def test_exact_line_spacing_does_not_leak_raw_emu(self):
        from docx import Document
        from docx.enum.text import WD_LINE_SPACING
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            p = doc.add_paragraph("Judul dengan spasi baris tetap.")
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(11.4)
            path = Path(tmp) / "exact_line_spacing.docx"
            doc.save(str(path))

            parser = DocxParser(path)
            para = next(pp for pp in parser.paragraphs if pp.text.startswith("Judul"))
            self.assertIsNone(para.line_spacing)

    def test_multiple_line_spacing_still_read_as_float(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            p = doc.add_paragraph("Paragraf spasi 1.5 normal.")
            p.paragraph_format.line_spacing = 1.5
            path = Path(tmp) / "multiple_line_spacing.docx"
            doc.save(str(path))

            parser = DocxParser(path)
            para = next(pp for pp in parser.paragraphs if pp.text.startswith("Paragraf"))
            self.assertAlmostEqual(para.line_spacing, 1.5)


if __name__ == "__main__":
    unittest.main(verbosity=2)
