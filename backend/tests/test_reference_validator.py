"""Unit tests untuk ReferenceValidator (sitasi in-text, DP, typo)."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest
from lxml import etree

from app.services.docx_parser import ParagraphInfo
from app.services.reference_validator import (
    ReferenceValidator,
    _acronym_author_match,
    _dp_has_dkk,
    _dp_has_et_al,
    _is_month_label,
)
from app.services.schema_rules import get_pkm_kc_proposal_rules


def _fake_parser(paragraphs: list[ParagraphInfo]) -> MagicMock:
    parser = MagicMock()
    parser.paragraphs = paragraphs
    parser.estimate_physical_page = None
    parser.estimate_paragraph_index_in_page = None

    def find_section_boundaries(
        section_names: list[str],
        case_sensitive: bool = False,
        headings_only: bool = False,
    ) -> dict[str, int | None]:
        result: dict[str, int | None] = {name: None for name in section_names}
        for para in paragraphs:
            if headings_only and not para.is_heading:
                continue
            text = para.text.strip()
            cmp_text = text if case_sensitive else text.upper()
            for name in section_names:
                if result[name] is not None:
                    continue
                cmp_name = name if case_sensitive else name.upper()
                if cmp_text.startswith(cmp_name):
                    result[name] = para.index
        return result

    parser.find_section_boundaries = find_section_boundaries
    return parser


def _heading(idx: int, text: str) -> ParagraphInfo:
    return ParagraphInfo(index=idx, text=text, is_heading=True)


def _body(idx: int, text: str) -> ParagraphInfo:
    return ParagraphInfo(index=idx, text=text, is_heading=False)


@pytest.fixture
def schema():
    return get_pkm_kc_proposal_rules()


def test_et_al_word_boundary_not_metal():
    assert _dp_has_et_al("metal oxide catalyst") is False
    assert _dp_has_et_al("Smith et al. (2020) Title.") is True
    assert _dp_has_et_al("Rahmadi et al (2022) paper") is True


def test_et_al_inside_reference_title_is_not_author_violation():
    """Regresi Egi/Fauzul: et al. setelah tahun berada di judul, bukan author."""
    entry = (
        "Fauzul, M.F., Hidayat, K., Ulya, M., Yaskun, M. dan Adhi, M. (2025) "
        "Analisis berdasarkan penelitian Ayun et al., 2020. Jurnal Contoh."
    )
    assert _dp_has_et_al(entry) is False


def test_dkk_detected():
    assert _dp_has_dkk("Penulis dkk. 2021.") is True
    assert _dp_has_dkk("No abbreviation here 2021.") is False


def test_dkk_inside_reference_title_is_not_author_violation():
    entry = (
        "Fauzul, M.F., Hidayat, K. dan Ulya, M. (2025) "
        "Perbandingan dengan penelitian Santoso dkk., 2023. Jurnal Contoh."
    )
    assert _dp_has_dkk(entry) is False


def test_intext_et_al_allowed(schema):
    """Sitasi in-text dengan 'et al.' DIPERBOLEHKAN — tidak boleh ke-flag."""
    paras = [
        _heading(0, "BAB 1. PENDAHULUAN"),
        _body(1, "Menurut (Revis et al, 2020) hal ini penting."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(
            3,
            "Revis, A. 2020. Judul buku yang cukup panjang agar tidak warning "
            "pendek. Penerbit: Kota.",
        ),
    ]
    parser = _fake_parser(paras)
    r = ReferenceValidator(parser, schema, current_year=2026).check()
    intext_fail = [
        f for f in r.format_issues if f.entry_index == -1 and "et al" in f.issue.lower()
    ]
    assert intext_fail == []


def test_intext_dkk_allowed(schema):
    """Sitasi in-text dengan 'dkk.' DIPERBOLEHKAN — tidak boleh ke-flag."""
    paras = [
        _heading(0, "BAB 1. X"),
        _body(1, "Lihat (Astuti dkk., 2012)."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(
            3,
            "Astuti. 2012. Judul artikel yang memadai untuk panjang minimum "
            "validator di sini. Jurnal Contoh, 1(1), 1–10.",
        ),
    ]
    parser = _fake_parser(paras)
    r = ReferenceValidator(parser, schema, current_year=2026).check()
    intext_fail = [
        f for f in r.format_issues if f.entry_index == -1 and "dkk" in f.issue.lower()
    ]
    assert intext_fail == []


def test_balance_typo_hint_same_year(schema):
    paras = [
        _heading(0, "BAB 1. X"),
        _body(1, "Menurut (Smyth, 2020) benar."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(
            3,
            "Smith, J. 2020. Solar Energy Fundamentals yang cukup panjang. "
            "New York: Springer.",
        ),
    ]
    parser = _fake_parser(paras)
    r = ReferenceValidator(parser, schema, current_year=2026).check()
    missing = [f for f in r.balance_findings if f.direction == "in_text_not_in_references"]
    assert missing
    assert "typo" in missing[0].detail.lower() or "smyth" in missing[0].detail.lower()


def test_dp_heading_but_no_entries_explains_gap(schema):
    """Judul DP ada, LAMPIRAN langsung setelah baris kosong → pesan spesifik."""
    paras = [
        _heading(0, "BAB 1. PENDAHULUAN"),
        _body(1, "Teks (Smith, 2020) dengan cukup panjang untuk pemeriksaan."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(3, ""),
        _heading(4, "LAMPIRAN"),
    ]
    parser = _fake_parser(paras)
    r = ReferenceValidator(parser, schema, current_year=2026).check()
    assert r.total_entries == 0
    assert r.dp_heading_paragraph_index == 2
    assert r.lampiran_heading_paragraph_index == 4
    joined = " ".join(m.text for m in r.messages)
    assert "terdeteksi" in joined.lower()
    assert "lampiran" in joined.lower()
    d = r.to_dict()
    assert d["section_detection"]["daftar_pustaka_heading_paragraph_index"] == 2


# ============================================================================
# Referensi terpotong paragraf kosong di tengah judul kutipan
# ============================================================================
#
# REGRESI lapangan (Leily, PKM-RSH artikel ilmiah): satu entry DP terpotong
# jadi dua paragraf Word oleh baris kosong di tengah judul —
#   "...Adolescents : A" | (baris kosong) | "Network Analysis’, Journal..."
# Potongan kedua tidak py tahun dan seharusnya digabung ke entry sebelumnya,
# tapi heuristik lama menghitung kata sebelum delimiter pertama ("Network
# Analysis’" → 2 kata) dan menyimpulkan itu entry baru (nama pengarang
# pendek) — padahal tanda kutip penutup (’) di situ adalah ekor judul
# artikel, bukan awal nama. Akibatnya DP terpecah jadi entry palsu yang lalu
# divonis "tahun tidak ditemukan" + "urutan alfabetis salah".


def test_split_reference_merged_via_trailing_title_quote(schema):
    paras = [
        _heading(0, "DAFTAR PUSTAKA"),
        _body(
            1,
            "Guadix, M.G., Sorrel, M.A. and Martínez, J. (2023) ‘Technology - "
            "Facilitated Sexual Violence Perpetration and Victimization Among "
            "Adolescents : A",
        ),
        _body(2, "Network Analysis’, Sexuality Research and Social Policy, 20(3), pp. 1000–1012."),
        _body(3, "Kahalon, R. and Klein, V. (2025) ‘Self-Objectification’, The Journal of Sex Research, 62(9), pp. 1722–1743."),
        _heading(4, "LAMPIRAN"),
    ]
    parser = _fake_parser(paras)
    r = ReferenceValidator(parser, schema, current_year=2026).check()
    assert r.total_entries == 2
    assert "Network Analysis" in r.entries[0].raw_text
    assert r.entries[0].year == 2023
    assert r.entries[1].author_first == "Kahalon"
    joined = " ".join(m.text for m in r.messages)
    assert "tahun" not in joined.lower() or "tidak ditemukan" not in joined.lower()


def test_split_reference_merged_via_page_range_continuation(schema):
    """Potongan kedua berupa 'pp. 1–12.' (nomor halaman) juga harus digabung."""
    paras = [
        _heading(0, "DAFTAR PUSTAKA"),
        _body(
            1,
            "Kang, H. (2021) ‘Sample size determination and power analysis’, "
            "Journal of Educational Evaluation for Health Professions,",
        ),
        _body(2, "pp. 1–12."),
        _heading(3, "LAMPIRAN"),
    ]
    parser = _fake_parser(paras)
    r = ReferenceValidator(parser, schema, current_year=2026).check()
    assert r.total_entries == 1
    assert r.entries[0].raw_text.endswith("pp. 1–12.")


# ============================================================================
# Bibliografi Mendeley berbungkus DUA lapis w:sdt
# ============================================================================
#
# REGRESI lapangan (progress-report PKM-RE): Mendeley terkadang membungkus
# field bibliografi dengan w:sdt bersarang (field wrapper luar + field
# "Bibliography" itu sendiri di dalamnya). `_walk_dp_body_siblings_ooxml`
# hanya mengambil <w:p> anak LANGSUNG dari sdtContent terluar — paragraf
# yang bersarang satu lapis lebih dalam (di sdtContent milik w:sdt kedua)
# tidak pernah kebaca. DAFTAR PUSTAKA lalu divonis kosong padahal Mendeley
# sudah mengisinya.


def _build_nested_sdt_bibliography_docx(out_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("DAFTAR PUSTAKA", level=1)
    body = doc.element.body
    # <w:sdt tag=MENDELEY_BIBLIOGRAPHY><w:sdtContent>
    #   <w:sdt tag=MENDELEY_BIBLIOGRAPHY><w:sdtContent>
    #     <w:p>entry 1</w:p>
    #     <w:p>entry 2</w:p>
    #   </w:sdtContent></w:sdt>
    # </w:sdtContent></w:sdt>
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    outer_xml = f"""
    <w:sdt xmlns:w="{ns}">
      <w:sdtPr><w:tag w:val="MENDELEY_BIBLIOGRAPHY"/></w:sdtPr>
      <w:sdtContent>
        <w:sdt>
          <w:sdtPr><w:tag w:val="MENDELEY_BIBLIOGRAPHY"/></w:sdtPr>
          <w:sdtContent>
            <w:p><w:r><w:t>Amin, A. (2022) 'Judul artikel pertama', Jurnal Contoh, 1(1), pp. 1-5.</w:t></w:r></w:p>
            <w:p><w:r><w:t>Budi, B. (2023) 'Judul artikel kedua', Jurnal Contoh, 2(1), pp. 6-10.</w:t></w:r></w:p>
          </w:sdtContent>
        </w:sdt>
      </w:sdtContent>
    </w:sdt>
    """
    sdt_el = etree.fromstring(outer_xml.encode("utf-8"))
    # body diakhiri <w:sectPr> — harus tetap jadi elemen terakhir, jadi
    # sisipkan SEBELUM itu (python-docx add_paragraph melakukan hal sama).
    from docx.oxml.ns import qn as _qn
    sect_pr = body.find(_qn("w:sectPr"))
    sect_pr.addprevious(sdt_el)
    doc.add_heading("LAMPIRAN", level=1)
    # add_heading juga menyisipkan sebelum sectPr, jadi urutannya sudah benar:
    # DAFTAR PUSTAKA, sdt, LAMPIRAN, sectPr.
    doc.save(str(out_path))


def test_nested_sdt_mendeley_bibliography_extracted(tmp_path, schema):
    from app.services.docx_parser import DocxParser

    out_path = tmp_path / "nested_sdt.docx"
    _build_nested_sdt_bibliography_docx(out_path)
    r = ReferenceValidator(DocxParser(out_path), schema, current_year=2026).check()
    assert r.total_entries == 2
    assert r.entries[0].author_first == "Amin"
    assert r.entries[1].author_first == "Budi"


NUKI_FIXTURE = Path(r"c:\Users\ac300\Downloads\PKM KC  NUKI OTISTA_evp_KP_030426.docx")


@pytest.mark.skipif(not NUKI_FIXTURE.exists(), reason="Lokal: file contoh Zotero tidak ada")
def test_zotero_sdt_bibliography_extracted(schema):
    """Dokumen dengan bibliografi di w:sdt (Zotero) punya entri DP terbaca."""
    from app.services.docx_parser import DocxParser

    r = ReferenceValidator(DocxParser(NUKI_FIXTURE), schema, current_year=2026).check()
    assert r.total_entries >= 10
    assert r.dp_heading_paragraph_index is not None
    assert r.entries[0].author_first is not None


# ============================================================================
# Keterangan bulan bukan sitasi
# ============================================================================
#
# Laporan kemajuan penuh keterangan jadwal yang bentuknya persis sitasi Harvard
# tanpa koma: "Tahap Monitoring dan Evaluasi (Agustus 2026)". Sebelum
# diperbaiki, itu divonis sitasi berformat salah, lalu ikut dicari di Daftar
# Pustaka dan dilaporkan lagi sebagai sitasi tak bersumber — satu keterangan
# jadwal menghasilkan dua temuan palsu.


@pytest.mark.parametrize(
    "teks", ["Agustus", "September", "Oktober", "Januari", "August", "Mei"]
)
def test_single_month_recognised(teks):
    assert _is_month_label(teks)


@pytest.mark.parametrize(
    "teks", ["Juli-Agustus", "Juli\u2013Agustus", "September\u2013Oktober"]
)
def test_month_range_recognised(teks):
    """Rentang dengan tanda hubung maupun en dash."""
    assert _is_month_label(teks)


@pytest.mark.parametrize(
    "teks", ["Sugiharti", "BPS", "Kotsios", "Hidayati dan Astuti", "Amin"]
)
def test_author_names_not_treated_as_month(teks):
    assert not _is_month_label(teks)


def test_schedule_heading_produces_no_citation():
    """REGRESI lapangan: sub-judul berketerangan bulan tidak boleh jadi sitasi."""
    parser = _fake_parser([
        ParagraphInfo(index=0, text="BAB 3. METODE PELAKSANAAN", is_heading=True),
        ParagraphInfo(
            index=1,
            text="Tahap Monitoring, Evaluasi, dan Keberlanjutan Program (Agustus 2026)",
        ),
        ParagraphInfo(index=2, text="Tahap Implementasi Program (Juli-Agustus 2026)"),
        ParagraphInfo(
            index=3,
            text="Kondisi ini dijelaskan penelitian sebelumnya (Kotsios, 2023).",
        ),
    ])
    v = ReferenceValidator(parser, get_pkm_kc_proposal_rules())
    citations = v._extract_in_text_citations()
    assert sorted(c.author for c in citations) == ["Kotsios"]


# ============================================================================
# Sitasi majemuk dipecah per sumber
# ============================================================================
#
# REGRESI lapangan (Lapkem PKM RSH): "(Hu et al., 2025; Tan dan Roswiyani,
# 2026)" terbaca satu sitasi "Hu, 2026" — penulis pertama dikawinkan dengan
# tahun terakhir — lalu dilaporkan tak ada di Daftar Pustaka padahal ada.


def _citations_of(*kalimat: str):
    paras = [ParagraphInfo(index=0, text="BAB 1. PENDAHULUAN", is_heading=True)]
    paras += [ParagraphInfo(index=i + 1, text=t) for i, t in enumerate(kalimat)]
    v = ReferenceValidator(_fake_parser(paras), get_pkm_kc_proposal_rules())
    return [(c.author, c.year, c.has_format_error) for c in v._extract_in_text_citations()]


def test_multi_citation_split_per_source():
    assert _citations_of(
        "Hal ini sejalan dengan temuan (Hu et al., 2025; Tan dan Roswiyani, 2026)."
    ) == [("Hu", 2025, False), ("Tan", 2026, False)]


def test_multi_citation_three_sources_no_space_after_semicolon():
    assert _citations_of(
        "Terbukti efektif (Sehgal et al., 2023;Júnior, 2022; Kong dkk., 2024)."
    ) == [("Sehgal", 2023, False), ("Júnior", 2022, False), ("Kong", 2024, False)]


def test_multi_citation_pieces_without_comma_still_format_error():
    """Tiap potongan dinilai aturan sitasi tunggal: tanpa koma = format salah."""
    assert _citations_of(
        "Metode ini umum dipakai (Johari et al. 2024;Salam et al. 2024)."
    ) == [("Johari", 2024, True), ("Salam", 2024, True)]


def test_multi_citation_non_breaking_space():
    assert _citations_of(
        "Dilaporkan sebelumnya (Sehgal\u00a0et al., 2023; Hu et al., 2025)."
    ) == [("Sehgal", 2023, False), ("Hu", 2025, False)]


def test_semicolon_without_two_years_not_split():
    """Kurung bertitik koma tanpa dua angka tahun bukan sitasi majemuk."""
    assert _citations_of(
        "Rincian ada di lampiran (lihat Tabel 1; Gambar 2).",
        "Data tahunan (2020-2021; 2022-2023) dirangkum.",
    ) == []


def test_prefix_match_tries_every_reference(schema):
    """REGRESI lapangan: dua instansi berawalan sama ("Kementrian ...").
    Sitasi singkatan harus dicocokkan ke entri yang tahunnya sesuai, bukan
    divonis beda tahun karena kandidat pertama kebetulan instansi lain."""
    paras = [
        _heading(0, "BAB 1. PENDAHULUAN"),
        _body(1, "Data UMKM (Kemenkop UKM, 2021) dan pariwisata (Kemenparekraf, 2020)."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(3, "Kementrian Koperasi dan Usaha Kecil dan Menengah Republik Indonesia. (2021) Perkembangan data usaha."),
        _body(4, "Kementrian Pariwisata dan Ekonomi Kreatif Republik Indonesia. (2020) Rencana strategis pengembangan."),
    ]
    r = ReferenceValidator(_fake_parser(paras), schema, current_year=2026).check()
    assert [f.detail for f in r.balance_findings if f.direction == "in_text_not_in_references"] == []


def test_prefix_match_still_flags_wrong_year(schema):
    paras = [
        _heading(0, "BAB 1. PENDAHULUAN"),
        _body(1, "Data pariwisata (Kemenparekraf, 2019)."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(3, "Kementrian Koperasi dan Usaha Kecil dan Menengah Republik Indonesia. (2021) Perkembangan data usaha."),
        _body(4, "Kementrian Pariwisata dan Ekonomi Kreatif Republik Indonesia. (2020) Rencana strategis pengembangan."),
    ]
    r = ReferenceValidator(_fake_parser(paras), schema, current_year=2026).check()
    missing = [f for f in r.balance_findings if f.direction == "in_text_not_in_references"]
    assert len(missing) == 1 and "2019" in missing[0].detail


@pytest.mark.parametrize("cite,dp", [
    ("Bappeda Kota Malang", "Badan Perencanaan Pembangunan Daerah Kota Malang"),
    ("BPS", "Badan Pusat Statistik"),
    ("Kemenkes RI", "Kementerian Kesehatan Republik Indonesia"),
    ("Kemenkop UKM", "Kementrian Koperasi dan Usaha Kecil dan Menengah Republik Indonesia"),
    ("Komdigi", "Kementerian Komunikasi dan Digital"),
    ("komdigi", "Kementerian Komunikasi dan Digital Republik Indonesia"),
    ("Kominfo", "Kementerian Komunikasi dan Informatika"),
])
def test_acronym_author_match(cite, dp):
    assert _acronym_author_match(cite, dp)


@pytest.mark.parametrize("cite,dp", [
    ("Bappeda Kota Batu", "Badan Perencanaan Pembangunan Daerah Kota Malang"),
    ("BPS", "Badan Pusat Statistik Jawa Timur"),
    ("Smith", "Smithson"),
    ("Sari", "Sekarsari"),
    ("Komdigi", "Kementerian Komunikasi dan Informatika"),
])
def test_acronym_author_mismatch(cite, dp):
    assert not _acronym_author_match(cite, dp)


def test_acronym_citation_matches_full_institution(schema):
    """REGRESI lapangan: "(Bappeda Kota Malang, 2025)" disitasi, DP menulis
    nama lengkap instansinya — sumbernya ada, bukan sitasi yatim."""
    paras = [
        _heading(0, "BAB 1. PENDAHULUAN"),
        _body(1, "Rencana kawasan (Bappeda Kota Malang, 2025) dan data lain (Bappeda Kota Malang, 2024)."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(3, "Badan Perencanaan Pembangunan Daerah Kota Malang. (2025) Rencana pengembangan kawasan heritage."),
    ]
    r = ReferenceValidator(_fake_parser(paras), schema, current_year=2026).check()
    missing = [f for f in r.balance_findings if f.direction == "in_text_not_in_references"]
    assert len(missing) == 1 and "2024" in missing[0].detail


def test_komdigi_citation_matches_full_institution(schema):
    """Nama pendek kementerian boleh menghilangkan awalan 'Kementerian'."""
    paras = [
        _heading(0, "BAB 1. PENDAHULUAN"),
        _body(1, "Transformasi digital meningkat (Komdigi, 2025)."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(
            3,
            "Kementerian Komunikasi dan Digital. (2025) Laporan transformasi "
            "digital nasional tahun 2025.",
        ),
    ]
    result = ReferenceValidator(_fake_parser(paras), schema, current_year=2026).check()

    missing = [
        finding
        for finding in result.balance_findings
        if finding.direction == "in_text_not_in_references"
    ]
    assert missing == []


def test_komdigi_citation_still_flags_different_year(schema):
    paras = [
        _heading(0, "BAB 1. PENDAHULUAN"),
        _body(1, "Transformasi digital meningkat (Komdigi, 2025)."),
        _heading(2, "DAFTAR PUSTAKA"),
        _body(
            3,
            "Kementerian Komunikasi dan Digital. (2024) Laporan transformasi "
            "digital nasional tahun 2024.",
        ),
    ]
    result = ReferenceValidator(_fake_parser(paras), schema, current_year=2026).check()

    missing = [
        finding
        for finding in result.balance_findings
        if finding.direction == "in_text_not_in_references"
    ]
    assert len(missing) == 1
    assert "bukan 2025" in missing[0].detail
