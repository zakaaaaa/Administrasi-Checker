"""
Generator varian PENYANDIAN dokumen — bahan uji invariansi.

Ide
---
Satu dokumen LOGIS yang sama (isi, struktur, urutan section identik) bisa
disandikan Word dengan banyak cara berbeda:

- Daftar Isi   : style bawaan 'toc 1' | dot leader manual | tab leader manual
- Heading      : style 'Heading N' | bold uppercase manual | style kustom
                 ber-outlineLvl
- Nomor halaman: part 'default' saja | part 'first' hidup (titlePg aktif) |
                 part 'first' mati (titlePg tidak aktif — Word tak merendernya)

Penempatan LOGIS nomor halaman selalu sama di ketiga gaya: romawi di bawah
untuk zona awal, arab di atas untuk zona inti. Yang berbeda hanya cara Word
menuliskannya di dalam file.

Semua varian ini WAJIB menghasilkan vonis yang sama. Kalau vonis berubah
padahal isinya tidak, itu bug — dan ketemu tanpa menunggu dokumen asli datang.

Ini melengkapi korpus dokumen nyata: korpus menjaga yang sudah pernah
diperbaiki tidak rusak lagi (regresi), varian menangkap bentuk penyandian yang
belum pernah dilihat (generalisasi) — sumber false positive yang selama ini
selalu muncul di dokumen baru.

Dipakai oleh tests/test_encoding_invariance.py.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# ---------------------------------------------------------------------------
# Isi logis — SATU sumber, dipakai semua varian
# ---------------------------------------------------------------------------

FRONT_MATTER = ["DAFTAR ISI", "DAFTAR LAMPIRAN"]

# (judul BAB, jumlah paragraf isi) — Laporan Kemajuan PKM-KC
CORE_SECTIONS = [
    ("BAB 1. PENDAHULUAN", 2),
    ("BAB 2. TARGET LUARAN", 2),
    ("BAB 3. TAHAP PELAKSANAAN", 2),
    ("BAB 4. HASIL YANG DICAPAI", 2),
    ("BAB 5. POTENSI HASIL", 2),
    ("BAB 6. RENCANA TAHAPAN BERIKUTNYA", 2),
]
CLOSING_SECTIONS = ["DAFTAR PUSTAKA", "LAMPIRAN"]

# Entri Daftar Isi + nomor halamannya (dipakai ketiga gaya TOC)
TOC_ROWS = [
    ("DAFTAR ISI", "i"),
    ("DAFTAR LAMPIRAN", "ii"),
    ("BAB 1. PENDAHULUAN", "1"),
    ("BAB 2. TARGET LUARAN", "2"),
    ("BAB 3. TAHAP PELAKSANAAN", "3"),
    ("BAB 4. HASIL YANG DICAPAI", "4"),
    ("BAB 5. POTENSI HASIL", "5"),
    ("BAB 6. RENCANA TAHAPAN BERIKUTNYA", "6"),
    ("DAFTAR PUSTAKA", "7"),
    ("LAMPIRAN", "8"),
]

BODY_TEXT = (
    "Kegiatan ini dilaksanakan sesuai rencana yang telah disusun pada tahap "
    "sebelumnya dengan memperhatikan ketercapaian target luaran."
)

TOC_STYLES = ("style", "dot_leader", "tab_leader")
HEADING_STYLES = ("heading_style", "manual_bold", "outline_level")
PAGENUM_STYLES = ("plain", "titlepg_live", "first_dead")


# ---------------------------------------------------------------------------
# Primitif penyandian
# ---------------------------------------------------------------------------


def _ensure_style(doc, name: str, builtin_base: str = "Normal"):
    """Buat paragraph style kustom kalau belum ada."""
    from docx.enum.style import WD_STYLE_TYPE

    if name in [s.name for s in doc.styles]:
        return doc.styles[name]
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[builtin_base]
    return st


def _add_toc_entry(doc, label: str, page: str, style: str) -> None:
    """Satu baris Daftar Isi dalam salah satu dari tiga gaya penyandian."""
    if style == "style":
        # Gaya Word asli: style bawaan 'toc 1', pemisah TAB.
        p = doc.add_paragraph(f"{label}\t{page}")
        try:
            p.style = doc.styles["toc 1"]
        except KeyError:
            _ensure_style(doc, "toc 1")
            p.style = doc.styles["toc 1"]
    elif style == "dot_leader":
        doc.add_paragraph(f"{label} {'.' * 20} {page}")
    elif style == "tab_leader":
        doc.add_paragraph(f"{label}\t{page}")
    else:
        raise ValueError(f"gaya TOC tidak dikenal: {style}")


def _add_heading(doc, text: str, style: str) -> None:
    """Satu heading section dalam salah satu dari tiga gaya penyandian."""
    if style == "heading_style":
        doc.add_heading(text, level=1)
    elif style == "manual_bold":
        # Style Normal, tapi teks bold — lazim di dokumen mahasiswa.
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
    elif style == "outline_level":
        # Style kustom yang namanya BUKAN "Heading N", ditandai outlineLvl.
        st = _ensure_style(doc, "JudulBab")
        p = doc.add_paragraph(text)
        p.style = st
        ppr = p._element.get_or_add_pPr()
        el = ppr.makeelement(qn("w:outlineLvl"), {qn("w:val"): "0"})
        ppr.append(el)
    else:
        raise ValueError(f"gaya heading tidak dikenal: {style}")


def _page_field(container, cached: str = "2") -> None:
    """Field PAGE rata kanan, TNR 12pt, lengkap dengan hasil ter-cache.

    Word selalu menyimpan hasil render terakhir sebagai <w:t> di antara
    fldChar begin/end. Tanpa itu checker menganggap part-nya tidak merender apa
    pun, sehingga alignment & font tidak ikut divalidasi.
    """
    p = container.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    el = run._element
    el.append(el.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"}))
    it = el.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
    it.text = "PAGE   \\* MERGEFORMAT"
    el.append(it)
    el.append(el.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "separate"}))
    t = el.makeelement(qn("w:t"), {})
    t.text = cached
    el.append(t)
    el.append(el.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"}))
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)


def _set_page_numbering(doc, style: str) -> None:
    """Penempatan LOGIS selalu sama — yang berbeda cuma cara menyandikannya.

    Aturan PKM: zona awal romawi di BAWAH, zona inti arab di ATAS. Ketiga gaya
    di bawah menyandikan penempatan yang sama persis, jadi vonisnya wajib sama.

    - plain        : footer 'default' (zona awal) + header 'default' (zona inti)
    - titlepg_live : sama, plus part 'first' berisi nomor DAN <w:titlePg/> aktif
                     → part 'first' benar-benar dirender Word
    - first_dead   : sama seperti plain, plus part 'first' berisi nomor TAPI
                     titlePg MATI → Word tidak pernah merendernya. Ini jebakan
                     yang ditemukan di dokumen PKM-RE asli; vonisnya harus
                     identik dengan 'plain'.
    """
    front, core = doc.sections[0], doc.sections[-1]

    front.footer.is_linked_to_previous = False
    _page_field(front.footer, cached="ii")
    core.header.is_linked_to_previous = False
    _page_field(core.header, cached="2")
    # Zona inti WAJIB memutus warisan footer zona depan. Tanpa ini footer
    # bernomor milik zona depan ikut terpakai di zona inti, sehingga nomor
    # halaman tercetak dua kali (atas dan bawah) — dokumen jadi tidak valid.
    # Di Word ini setara: masuk ke footer bagian inti, matikan "Link to
    # Previous", lalu kosongkan isinya.
    core.footer.is_linked_to_previous = False

    if style == "plain":
        return

    if style == "titlepg_live":
        # Zona inti juga harus memutus warisan footer halaman-pertama milik
        # zona depan — kalau tidak, halaman pertama zona inti mencetak nomor
        # di header (miliknya) DAN footer (warisan) sekaligus.
        core.different_first_page_header_footer = True
        core.first_page_footer.is_linked_to_previous = False
        for sec, container_attr, cached in (
            (front, "first_page_footer", "i"),
            (core, "first_page_header", "1"),
        ):
            sec.different_first_page_header_footer = True
            container = getattr(sec, container_attr)
            container.is_linked_to_previous = False
            _page_field(container, cached=cached)
        return

    if style == "first_dead":
        # Bikin part 'first' dulu, lalu MATIKAN titlePg-nya. Referensinya tetap
        # tertulis di sectPr — persis kondisi dokumen PKM-RE.
        front.different_first_page_header_footer = True
        front.first_page_header.is_linked_to_previous = False
        _page_field(front.first_page_header, cached="i")
        front.different_first_page_header_footer = False
        return

    raise ValueError(f"gaya nomor halaman tidak dikenal: {style}")


def _set_numeral_format(section, fmt: str, start: int | None = None) -> None:
    """<w:pgNumType w:fmt="..."/> pada sectPr."""
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    attrs = {qn("w:fmt"): fmt}
    if start is not None:
        attrs[qn("w:start")] = str(start)
    sect_pr.append(sect_pr.makeelement(qn("w:pgNumType"), attrs))


# ---------------------------------------------------------------------------
# Builder
# ---------------------------------------------------------------------------


def build_variant(
    out_path: Path,
    *,
    toc_style: str = "style",
    heading_style: str = "heading_style",
    pagenum_style: str = "plain",
) -> Path:
    """Bangun satu varian penyandian dari dokumen logis yang sama."""
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    first = doc.sections[0]
    first.left_margin = Cm(4)
    first.right_margin = Cm(3)
    first.top_margin = Cm(3)
    first.bottom_margin = Cm(3)
    _set_numeral_format(first, "lowerRoman", start=1)

    # ---- Zona awal ----
    _add_heading(doc, "DAFTAR ISI", heading_style)
    for label, page in TOC_ROWS:
        _add_toc_entry(doc, label, page, toc_style)

    _add_heading(doc, "DAFTAR LAMPIRAN", heading_style)
    doc.add_paragraph("Lampiran 1. Penggunaan Dana")
    doc.add_paragraph("Lampiran 2. Bukti Pendukung Kegiatan")

    # ---- Zona inti (section baru, penomoran arab) ----
    core = doc.add_section(WD_SECTION.NEW_PAGE)
    core.left_margin = Cm(4)
    core.right_margin = Cm(3)
    core.top_margin = Cm(3)
    core.bottom_margin = Cm(3)
    _set_numeral_format(core, "decimal", start=1)

    for title, n_body in CORE_SECTIONS:
        _add_heading(doc, title, heading_style)
        for _ in range(n_body):
            doc.add_paragraph(BODY_TEXT)

    _add_heading(doc, "DAFTAR PUSTAKA", heading_style)
    doc.add_paragraph(
        "Makoy, C. G. I. (2026) Sintesis komposit untuk fotokatalisis. "
        "Jurnal Kimia Indonesia, 12(2), pp. 45-58."
    )

    _add_heading(doc, "LAMPIRAN", heading_style)
    doc.add_paragraph("Lampiran 1. Penggunaan Dana")
    doc.add_paragraph("Lampiran 2. Bukti Pendukung Kegiatan")

    _set_page_numbering(doc, pagenum_style)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def build_all_variants(out_dir: Path) -> dict[tuple[str, str, str], Path]:
    """Bangun seluruh kombinasi varian. Return {(toc, heading, pagenum): path}."""
    built: dict[tuple[str, str, str], Path] = {}
    for toc in TOC_STYLES:
        for head in HEADING_STYLES:
            for pnum in PAGENUM_STYLES:
                key = (toc, head, pnum)
                path = out_dir / f"variant_{toc}__{head}__{pnum}.docx"
                build_variant(
                    path, toc_style=toc, heading_style=head, pagenum_style=pnum
                )
                built[key] = path
    return built


if __name__ == "__main__":
    import sys

    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("variants")
    made = build_all_variants(out)
    print(f"{len(made)} varian dibuat di {out}/")
