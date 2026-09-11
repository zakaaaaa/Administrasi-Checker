"""Pengecualian temuan author hilang untuk nama seluruhnya kapital."""

import pytest
from docx import Document

from app.services.docx_parser import DocxParser
from app.services.reference_validator import ReferenceValidator
from app.services.schema_rules import get_pkm_kc_proposal_rules


@pytest.fixture
def check_document(tmp_path):
    def check(text, reference_author="Sari", reference_year=2020):
        doc = Document()
        doc.add_heading("BAB 1. PENDAHULUAN", level=1)
        doc.add_paragraph(text)
        doc.add_heading("DAFTAR PUSTAKA", level=1)
        doc.add_paragraph(
            f"{reference_author}. {reference_year}. Judul referensi contoh "
            "untuk pemeriksaan sitasi dalam laporan."
        )
        path = tmp_path / "uppercase_author.docx"
        doc.save(path)
        return ReferenceValidator(
            DocxParser(path),
            get_pkm_kc_proposal_rules(),
            minimum_recommended_recent=0,
        ).check()

    return check


@pytest.mark.parametrize("author", ["BPS", "WHO", "PPATK", "BPS RI", "KEMENTERIAN KESEHATAN"])
def test_uppercase_missing_author_is_kept_as_citation_without_failure(check_document, author):
    result = check_document(f"Data bersumber dari ({author}, 2023).")

    assert [(c.author, c.year) for c in result.in_text_citations] == [(author, 2023)]
    assert result.balance_findings == []
    assert result.to_dict()["citation_balance"]["in_text_not_in_references"] == []
    assert not any("tidak ditemukan" in m.text or "tidak ada author" in m.text for m in result.messages)
    assert result.status == "pass"


@pytest.mark.parametrize("author", ["Bps", "bps", "Pratama", "123"])
def test_other_missing_authors_still_reported(check_document, author):
    result = check_document(f"Data bersumber dari ({author}, 2023).")

    assert len(result.balance_findings) == 1
    assert result.balance_findings[0].citation_or_entry == f"({author}, 2023)"
    assert result.status == "fail"


def test_uppercase_author_with_different_reference_year_still_fails(check_document):
    result = check_document("Data (BPS, 2023).", reference_author="BPS", reference_year=2022)

    assert len(result.balance_findings) == 1
    assert "bukan 2023" in result.balance_findings[0].detail
    assert result.status == "fail"


def test_uppercase_missing_author_does_not_hide_missing_comma(check_document):
    result = check_document("Data (BPS 2023).")

    assert len(result.in_text_citations) == 1
    assert result.balance_findings == []
    assert len(result.format_issues) == 1
    assert result.format_issues[0].entry_index == -1
    assert "koma" in result.format_issues[0].issue
    assert result.status == "fail"


def test_uppercase_exception_does_not_hide_other_missing_source(check_document):
    result = check_document("Data (BPS, 2023) dibandingkan dengan (Pratama, 2023).")

    assert len(result.in_text_citations) == 2
    assert len(result.balance_findings) == 1
    assert result.balance_findings[0].citation_or_entry == "(Pratama, 2023)"
    assert any(m.text == "1 sitasi di teks tidak ditemukan di Daftar Pustaka." for m in result.messages)
    assert not any("'BPS'" in m.text for m in result.messages)
